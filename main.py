import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

FILE_NAME = "uriage.xlsx"

# 商品マスタ
PRODUCTS = {
    "4901234567890": {"name": "ジュース", "price": 500},
    "4901234567891": {"name": "お茶", "price": 300},
    "4901234567892": {"name": "スポーツドリンク", "price": 200},
}

# Excel初期化
if not os.path.exists(FILE_NAME):
    wb = Workbook()

    ws_sales = wb.active
    ws_sales.title = "売上データ"
    ws_sales.append(
        ["時刻", "商品名", "本数", "金額", "番号", "状態"]
    )

    ws_stock = wb.create_sheet("在庫データ")
    ws_stock.append(["在庫数"])
    ws_stock.append([0])

    wb.save(FILE_NAME)

order_number = 1


def get_stock():
    wb = load_workbook(FILE_NAME)
    ws = wb["在庫データ"]
    return ws["A2"].value


def update_stock(change):
    wb = load_workbook(FILE_NAME)
    ws = wb["在庫データ"]

    current_stock = ws["A2"].value

    if current_stock is None:
        current_stock = 0

    ws["A2"] = current_stock + change

    wb.save(FILE_NAME)


def refresh_stock_label():
    stock_var.set(f"現在の在庫 : {get_stock()} 本")


def update_excel_state(order_no, state):

    wb = load_workbook(FILE_NAME)
    ws = wb["売上データ"]

    for row in ws.iter_rows(min_row=2):
        if row[4].value == int(order_no):
            row[5].value = state
            break

    wb.save(FILE_NAME)


def save_sale(product_name, price):

    global order_number

    if get_stock() <= 0:
        messagebox.showerror(
            "在庫不足",
            "在庫がありません。"
        )
        return

    now = datetime.now().strftime("%H:%M:%S")

    wb = load_workbook(FILE_NAME)
    ws = wb["売上データ"]

    ws.append([
        now,
        product_name,
        1,
        price,
        order_number,
        "有効"
    ])

    wb.save(FILE_NAME)

    update_stock(-1)

    show_sales()

    refresh_stock_label()

    order_number += 1


def barcode_scan(event=None):

    code = barcode_entry.get().strip()

    if code not in PRODUCTS:

        messagebox.showerror(
            "エラー",
            f"未登録バーコード\n{code}"
        )

        barcode_entry.delete(0, tk.END)
        return

    product = PRODUCTS[code]

    save_sale(
        product["name"],
        product["price"]
    )

    barcode_entry.delete(0, tk.END)


def save_stock():

    qty = stock_entry.get()

    if not qty:
        return

    update_stock(int(qty))

    refresh_stock_label()

    messagebox.showinfo(
        "完了",
        f"{qty}本追加しました"
    )

    stock_entry.delete(0, tk.END)


def show_sales():

    for item in sales_tree.get_children():
        sales_tree.delete(item)

    wb = load_workbook(FILE_NAME)
    ws = wb["売上データ"]

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        check = "☑" if row[5] == "無効" else "☐"

        iid = sales_tree.insert(
            "",
            tk.END,
            values=(check, *row)
        )

        if row[5] == "無効":
            sales_tree.item(
                iid,
                tags=("disabled",)
            )


def toggle_check(event):

    item = sales_tree.identify_row(event.y)

    if not item:
        return

    values = list(
        sales_tree.item(item, "values")
    )

    if values[0] == "☐":

        values[0] = "☑"
        values[-1] = "無効"

        sales_tree.item(
            item,
            tags=("disabled",)
        )

    else:

        values[0] = "☐"
        values[-1] = "有効"

        sales_tree.item(
            item,
            tags=()
        )

    sales_tree.item(
        item,
        values=values
    )

    update_excel_state(
        values[5],
        values[-1]
    )


root = tk.Tk()
root.title("バーコード売上管理システム")
root.geometry("1000x650")

stock_var = tk.StringVar()

stock_label = ttk.Label(
    root,
    textvariable=stock_var,
    font=("Arial", 16)
)
stock_label.pack(pady=10)

refresh_stock_label()

notebook = ttk.Notebook(root)
notebook.pack(
    fill="both",
    expand=True
)

# --------------------
# 売上画面
# --------------------

frame_sales = ttk.Frame(notebook)
notebook.add(
    frame_sales,
    text="売上入力"
)

input_frame = ttk.LabelFrame(
    frame_sales,
    text="バーコード入力",
    padding=10
)
input_frame.pack(
    side="left",
    fill="y",
    padx=10,
    pady=10
)

ttk.Label(
    input_frame,
    text="バーコード"
).grid(row=0, column=0)

barcode_entry = ttk.Entry(
    input_frame,
    width=30
)
barcode_entry.grid(
    row=0,
    column=1
)

barcode_entry.bind(
    "<Return>",
    barcode_scan
)

barcode_entry.focus()

ttk.Label(
    input_frame,
    text="商品を読み取るだけで販売されます"
).grid(
    row=1,
    columnspan=2,
    pady=10
)

list_frame = ttk.LabelFrame(
    frame_sales,
    text="売上一覧",
    padding=10
)
list_frame.pack(
    side="right",
    fill="both",
    expand=True,
    padx=10,
    pady=10
)

columns = (
    "選択",
    "時刻",
    "商品名",
    "本数",
    "金額",
    "番号",
    "状態",
)

sales_tree = ttk.Treeview(
    list_frame,
    columns=columns,
    show="headings"
)

for col in columns:
    sales_tree.heading(col, text=col)
    sales_tree.column(
        col,
        width=100,
        anchor="center"
    )

sales_tree.pack(
    fill="both",
    expand=True
)

sales_tree.bind(
    "<Button-1>",
    toggle_check
)

sales_tree.tag_configure(
    "disabled",
    foreground="gray"
)

# --------------------
# 仕入れ画面
# --------------------

frame_stock = ttk.Frame(notebook)

notebook.add(
    frame_stock,
    text="仕入れ入力"
)

stock_frame = ttk.LabelFrame(
    frame_stock,
    text="仕入れ",
    padding=10
)

stock_frame.pack(
    padx=10,
    pady=10
)

ttk.Label(
    stock_frame,
    text="追加本数"
).grid(row=0, column=0)

stock_entry = ttk.Entry(
    stock_frame
)

stock_entry.grid(
    row=0,
    column=1
)

ttk.Button(
    stock_frame,
    text="追加",
    command=save_stock
).grid(
    row=1,
    columnspan=2,
    pady=10
)

show_sales()

root.mainloop()
